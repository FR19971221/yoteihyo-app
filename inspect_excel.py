import openpyxl
import json
import sys

def analyze():
    wb = openpyxl.load_workbook('東京0819予定表.xlsx', data_only=True)
    print("Sheets:", wb.sheetnames)
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        print(f"\n=== Sheet: {sheetname} ({sheet.max_row} rows, {sheet.max_column} cols) ===")
        # Look for headers or dates
        for r in range(1, min(25, sheet.max_row + 1)):
            row_str = []
            for c in range(1, min(35, sheet.max_column + 1)):
                val = sheet.cell(row=r, column=c).value
                if val is not None:
                    row_str.append(f"({r},{c}):{str(val).strip()[:20]}")
            if row_str:
                print(f"Row {r:2d}: " + " | ".join(row_str[:8]))

if __name__ == '__main__':
    analyze()
