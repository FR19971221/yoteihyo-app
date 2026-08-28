import openpyxl
import json
import sys

def inspect_sheet():
    wb = openpyxl.load_workbook('東京0819予定表.xlsx', data_only=True)
    print("Sheets:", wb.sheetnames)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n--- Sheet: {sname} ({ws.max_row} rows, {ws.max_column} cols) ---")
        for r in range(1, min(40, ws.max_row + 1)):
            row_items = []
            for c in range(1, min(40, ws.max_column + 1)):
                cell = ws.cell(row=r, column=c)
                v = cell.value
                # check borders
                borders = []
                if cell.border.left and cell.border.left.style: borders.append("L")
                if cell.border.right and cell.border.right.style: borders.append("R")
                if cell.border.top and cell.border.top.style: borders.append("T")
                if cell.border.bottom and cell.border.bottom.style: borders.append("B")
                b_str = "".join(borders)
                if v is not None or b_str:
                    txt = str(v).replace('\n', ' ')[:15] if v is not None else ""
                    row_items.append(f"C{c}[{b_str}]:{txt}")
            if row_items:
                print(f"R{r:2d}: " + " | ".join(row_items[:10]))

if __name__ == '__main__':
    inspect_sheet()
